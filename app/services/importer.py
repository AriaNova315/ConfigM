"""
Excel 导入/导出服务。
"""
import re
import io
import json
from datetime import datetime

import openpyxl
from openpyxl import Workbook
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.models import (
    Project, Phase, Activity, ActivityType, ActivityProduct,
    Product, Baseline, BaselineProduct,
)
from app.services.sequence import assign_numbers_for_activity

# 模块级临时存储（单用户应用，parse → confirm 之间保存解析结果）
_pending: dict = {}


# ── 解析工具 ──────────────────────────────────────────────

def _parse_date(val) -> str | None:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    return s or None


def _wbs_level(wbs) -> int:
    s = str(wbs).strip()
    if re.match(r'^\d+$', s):
        return 1
    return s.count('.') + 1


def _activity_type_name(task: str) -> str | None:
    if '入库' in task:
        return '产品入库'
    if '出库' in task:
        return '产品出库'
    if '发布' in task and '基线' in task:
        return '发布基线'
    if '配置状态纪实' in task:
        return '配置状态纪实'
    if '配置管理报告' in task:
        return '配置管理报告'
    return None


def _match_baseline(task: str, bl_names: list[str]) -> str | None:
    for name in bl_names:
        if name in task:
            return name
    return None


# ── 解析 Excel ────────────────────────────────────────────

def parse_excel(file_bytes: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    # 项目信息
    ws = wb['项目信息']
    r = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
    project = {
        'name': str(r[0] or '').strip(),
        'code': str(r[1] or '').strip(),
        'version': '0.0',
        'leader': str(r[2] or '').strip(),
        'cm_manager': str(r[3] or '').strip(),
    }

    # 配置项清单
    ws = wb['配置项清单']
    products = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[1]:
            continue
        products.append({
            'name': str(row[1]).strip(),
            'code': str(row[2] or '').strip(),
            'version': str(row[3] or 'V1.0').strip(),
        })
    product_by_name = {p['name']: p for p in products}

    # 基线组成
    ws = wb['基线组成']
    baselines = []
    cur_bl = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            bl_code = str(row[1] or '').strip()
            bl_version = str(row[2] or '').strip() or 'V1.0'
            cur_bl = {
                'name': str(row[0]).strip(),
                'code': bl_code,
                'version': bl_version,
                'status': '有效',
                'items': [],
            }
            baselines.append(cur_bl)
        if cur_bl and row[3]:
            cur_bl['items'].append({
                'product_name': str(row[3]).strip(),
                'product_code': str(row[4] or '').strip(),
                'product_version': str(row[5] or 'V1.0').strip(),
            })
    bl_names = [b['name'] for b in baselines]

    # 项目进度计划
    ws = wb['项目进度计划']
    phases = []
    cur_phase = None
    cur_cat = None

    for row in ws.iter_rows(min_row=2, values_only=True):
        wbs, task = row[0], str(row[1] or '').strip()
        start, end = _parse_date(row[2]), _parse_date(row[3])
        related = str(row[4] or '').strip()

        if wbs is None or not task:
            continue

        level = _wbs_level(wbs)

        if level == 1:
            cur_phase = {
                'name': task,
                'start_date': None,
                'end_date': None,
                '_starts': [],
                '_ends': [],
                'activities': [],
            }
            phases.append(cur_phase)
            cur_cat = None

        elif level == 2:
            cur_cat = task

        elif level == 3:
            if cur_phase:
                if start: cur_phase['_starts'].append(start)
                if end:   cur_phase['_ends'].append(end)

            if cur_cat == '配置管理活动' and cur_phase:
                at_name = _activity_type_name(task)
                if not at_name:
                    continue
                product_names = (
                    [n.strip() for n in re.split('[，,]', related) if n.strip()]
                    if related else []
                )
                cur_phase['activities'].append({
                    'task_name': task,
                    'activity_type_name': at_name,
                    'activity_date': end or start or '',
                    'product_names': product_names,
                    'matched_baseline': _match_baseline(task, bl_names) if at_name == '发布基线' else None,
                    'checkin_reason': None,
                    'checkout_reason': None,
                })

    for ph in phases:
        if ph['_starts']: ph['start_date'] = min(ph['_starts'])
        if ph['_ends']:   ph['end_date']   = max(ph['_ends'])

    # 需要确认的活动
    needs_checkin, needs_checkout, needs_baseline = [], [], []
    for ph in phases:
        for act in ph['activities']:
            key = f"{ph['name']}||{act['task_name']}||{act['activity_date']}"
            entry = {
                'key': key,
                'phase_name': ph['name'],
                'task_name': act['task_name'],
                'activity_date': act['activity_date'],
                'product_names': act['product_names'],
            }
            if act['activity_type_name'] == '产品入库':
                needs_checkin.append(entry)
            elif act['activity_type_name'] == '产品出库':
                needs_checkout.append(entry)
            elif act['activity_type_name'] == '发布基线':
                needs_baseline.append({**entry, 'auto_matched': act['matched_baseline']})

    return {
        'project': project,
        'products': products,
        'baselines': baselines,
        'phases': phases,
        'needs_checkin': needs_checkin,
        'needs_checkout': needs_checkout,
        'needs_baseline': needs_baseline,
    }


# ── 执行导入 ──────────────────────────────────────────────

async def clear_db(db: AsyncSession):
    for tbl in ['document_records', 'activity_products', 'baseline_products',
                'activities', 'baselines', 'products', 'phases', 'projects']:
        await db.execute(text(f'DELETE FROM {tbl}'))
    await db.flush()


async def execute_import(
    db: AsyncSession,
    data: dict,
    checkin_reasons: dict[str, str],
    checkout_reasons: dict[str, str],
    baseline_selections: dict[str, str] | None = None,
) -> dict:
    await clear_db(db)

    at_rows = await db.execute(select(ActivityType))
    at_map: dict[str, ActivityType] = {a.name: a for a in at_rows.scalars().all()}

    # 项目
    proj = Project(**data['project'])
    db.add(proj)
    await db.flush()

    # 配置项
    # prod_map:        名称 → Product（活动关联回退用，同名多版本时指向最后一条）
    # prod_nv_map:     (名称, 版本) → Product（基线关联精确匹配用）
    # prod_concat_map: 名称+版本拼接字符串 → Product（Excel 关联列写法，如"源代码V0.0.3"）
    prod_map: dict[str, Product] = {}
    prod_nv_map: dict[tuple, Product] = {}
    prod_concat_map: dict[str, Product] = {}
    for pd in data['products']:
        p = Product(project_id=proj.id, name=pd['name'], code=pd['code'],
                    version=pd['version'], product_type='文档')
        db.add(p)
        await db.flush()
        prod_map[pd['name']] = p
        prod_nv_map[(pd['name'], pd['version'])] = p
        prod_concat_map[pd['name'] + pd['version']] = p

    # 基线
    bl_map: dict[str, Baseline] = {}
    for bd in data['baselines']:
        bl = Baseline(project_id=proj.id, name=bd['name'], code=bd['code'],
                      version=bd['version'], status=bd['status'])
        db.add(bl)
        await db.flush()
        bl_map[bd['name']] = bl
        for i, item in enumerate(bd['items']):
            prod = prod_nv_map.get((item['product_name'], item['product_version'])) \
                   or prod_map.get(item['product_name'])
            if prod:
                db.add(BaselineProduct(baseline_id=bl.id, product_id=prod.id, sort_order=i))
        await db.flush()

    # 阶段和活动
    total_activities = 0
    for ph_data in data['phases']:
        if not ph_data['activities']:
            continue
        phase = Phase(project_id=proj.id, name=ph_data['name'],
                      start_date=ph_data.get('start_date'),
                      end_date=ph_data.get('end_date'))
        db.add(phase)
        await db.flush()

        dt_cnt: dict[tuple, int] = {}
        for act_data in ph_data['activities']:
            at = at_map.get(act_data['activity_type_name'])
            if not at:
                continue

            key = f"{ph_data['name']}||{act_data['task_name']}||{act_data['activity_date']}"
            checkin_reason = checkin_reasons.get(key) if act_data['activity_type_name'] == '产品入库' else None
            checkout_reason = checkout_reasons.get(key) if act_data['activity_type_name'] == '产品出库' else None

            baseline_id = None
            if act_data['activity_type_name'] == '发布基线':
                key = f"{ph_data['name']}||{act_data['task_name']}||{act_data['activity_date']}"
                bl_name = (baseline_selections or {}).get(key)
                bl = bl_map.get(bl_name) if bl_name else None
                if bl:
                    baseline_id = bl.id
                    bl.release_date = act_data['activity_date']

            dt_key = (act_data['activity_date'], at.id)
            dt_cnt[dt_key] = dt_cnt.get(dt_key, 0) + 1

            activity = Activity(
                phase_id=phase.id,
                activity_type_id=at.id,
                activity_date=act_data['activity_date'],
                notes=act_data['task_name'],
                checkin_reason=checkin_reason,
                baseline_id=baseline_id,
                checkout_reason=checkout_reason,
                entry_order=dt_cnt[dt_key],
            )
            db.add(activity)
            await db.flush()

            for i, pname in enumerate(act_data['product_names']):
                prod = prod_concat_map.get(pname) or prod_map.get(pname)
                if prod:
                    db.add(ActivityProduct(activity_id=activity.id, product_id=prod.id, sort_order=i))
            await db.flush()

            await assign_numbers_for_activity(db, activity, proj.id, proj.code, proj.version)
            total_activities += 1

    await _apply_version_status(db, proj.id)
    await db.commit()
    return {
        'project': data['project']['name'],
        'phases': len([p for p in data['phases'] if p['activities']]),
        'products': len(data['products']),
        'baselines': len(data['baselines']),
        'activities': total_activities,
    }


# ── 版本状态后处理 ─────────────────────────────────────────

async def _apply_version_status(db: AsyncSession, project_id: int):
    """
    同名同标识的多版本配置项：按入库活动日期升序排列，
    最晚入库的版本为有效，其余为无效；并设置 previous_version。
    """
    # 查询所有产品及其最早入库日期
    rows = await db.execute(
        select(Product, Activity.activity_date)
        .join(ActivityProduct, ActivityProduct.product_id == Product.id)
        .join(Activity, Activity.id == ActivityProduct.activity_id)
        .join(ActivityType, ActivityType.id == Activity.activity_type_id)
        .where(Product.project_id == project_id)
        .where(ActivityType.name == '产品入库')
        .order_by(Activity.activity_date, Activity.entry_order)
    )
    # 按 (name, code) 分组，保留每个 product_id 最早的入库日期
    from collections import defaultdict
    group: dict[tuple, list[tuple]] = defaultdict(list)  # (name,code) → [(date, product)]
    seen: set[int] = set()
    for prod, date in rows.all():
        if prod.id not in seen:
            seen.add(prod.id)
            group[(prod.name, prod.code)].append((date, prod))

    for versions in group.values():
        if len(versions) <= 1:
            continue
        # 按入库日期升序
        versions.sort(key=lambda x: x[0])
        prev_version = None
        for i, (_, prod) in enumerate(versions):
            prod.previous_version = prev_version
            prod.status = '有效' if i == len(versions) - 1 else '无效'
            prev_version = prod.version
    await db.flush()


# ── 导出 Excel ────────────────────────────────────────────

async def build_export(db: AsyncSession) -> bytes:
    from sqlalchemy.orm import selectinload
    from app.models.models import DocumentRecord

    proj_row = await db.execute(select(Project))
    proj = proj_row.scalars().first()
    if not proj:
        return b''

    # 配置项
    prod_rows = await db.execute(
        select(Product).where(Product.project_id == proj.id).order_by(Product.id)
    )
    products = prod_rows.scalars().all()
    prod_by_id = {p.id: p for p in products}

    # 基线
    bl_rows = await db.execute(
        select(Baseline).where(Baseline.project_id == proj.id).order_by(Baseline.id)
    )
    baselines = bl_rows.scalars().all()

    bp_rows = await db.execute(
        select(BaselineProduct).order_by(BaselineProduct.baseline_id, BaselineProduct.sort_order)
    )
    bp_list = bp_rows.scalars().all()
    bl_items: dict[int, list] = {}
    for bp in bp_list:
        bl_items.setdefault(bp.baseline_id, []).append(bp)

    # 阶段和活动
    ph_rows = await db.execute(
        select(Phase).where(Phase.project_id == proj.id).order_by(Phase.id)
    )
    phases = ph_rows.scalars().all()

    phase_acts: dict[int, list] = {}
    for ph in phases:
        act_rows = await db.execute(
            select(Activity).where(Activity.phase_id == ph.id)
            .order_by(Activity.activity_date, Activity.entry_order)
        )
        acts = act_rows.scalars().all()
        act_list = []
        for act in acts:
            ap_rows = await db.execute(
                select(ActivityProduct).where(ActivityProduct.activity_id == act.id)
                .order_by(ActivityProduct.sort_order)
            )
            ap_list = ap_rows.scalars().all()
            at = await db.get(ActivityType, act.activity_type_id)
            act_list.append({
                'task_name': act.notes or (at.name if at else ''),
                'activity_date': act.activity_date,
                'product_names': [prod_by_id[ap.product_id].name for ap in ap_list if ap.product_id in prod_by_id],
            })
        phase_acts[ph.id] = act_list

    # 构建 Excel
    wb = Workbook()

    ws1 = wb.active
    ws1.title = '项目信息'
    ws1.append(['项目名称', '项目代号', '项目组长', '配置管理员'])
    ws1.append([proj.name, proj.code, proj.leader, proj.cm_manager])

    ws2 = wb.create_sheet('项目进度计划')
    ws2.append(['WBS', '任务名称', '计划开始时间', '计划完成时间', '关联配置项'])
    for i, ph in enumerate(phases, start=1):
        acts = phase_acts.get(ph.id, [])
        if not acts:
            continue
        ws2.append([i, ph.name, ph.start_date or '', ph.end_date or '', None])
        ws2.append([f'{i}.1', '配置管理活动', None, None, None])
        for j, act in enumerate(acts, start=1):
            ws2.append([
                f'{i}.1.{j}',
                act['task_name'],
                act['activity_date'],
                act['activity_date'],
                '，'.join(act['product_names']) or None,
            ])

    ws3 = wb.create_sheet('配置项清单')
    ws3.append(['序号', '配置项名称', '配置项标识', '版本'])
    for i, p in enumerate(products, start=1):
        ws3.append([i, p.name, p.code, p.version])

    ws4 = wb.create_sheet('基线组成')
    ws4.append(['基线名称', '基线标识', '基线版本', '包含配置项名称', '配置项标识', '配置项版本'])
    for bl in baselines:
        items = bl_items.get(bl.id, [])
        for j, bp in enumerate(items):
            prod = prod_by_id.get(bp.product_id)
            ws4.append([
                bl.name if j == 0 else None,
                bl.code if j == 0 else None,
                bl.version if j == 0 else None,
                prod.name if prod else '',
                prod.code if prod else '',
                prod.version if prod else '',
            ])
        if not items:
            ws4.append([bl.name, bl.code, bl.version, None, None, None])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
