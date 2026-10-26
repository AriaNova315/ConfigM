"""
build_templates.py
用 python-docx 将原始 Word 模板转换为 docxtpl 占位符模板。
保留原始格式（字体、边框、合并单元格），只替换文字内容。
输出到 templates/tpl/*_tpl.docx
"""

import copy
import os
from docx import Document
from docx.oxml.ns import qn

SRC = os.path.join(os.path.dirname(__file__), "templates")
DST = os.path.join(os.path.dirname(__file__), "templates", "tpl")
os.makedirs(DST, exist_ok=True)


# ── 辅助函数 ──────────────────────────────────────────────────────────────

def set_para(para, text):
    """清空段落所有 run，写入单条 run（保留第一个 run 的格式）。"""
    if not para.runs:
        para.add_run(text)
        return
    first_run = para.runs[0]
    rpr = first_run._r.find(qn("w:rPr"))
    for run in para.runs:
        run._r.getparent().remove(run._r)
    new_run = para.add_run(text)
    if rpr is not None:
        new_run._r.insert(0, copy.deepcopy(rpr))


def set_cell(cell, text):
    """清空单元格所有段落内容，在第一个段落写入单条 run。"""
    for i, para in enumerate(cell.paragraphs):
        for run in para.runs:
            run._r.getparent().remove(run._r)
    para = cell.paragraphs[0]
    if para.runs:
        for run in para.runs:
            run._r.getparent().remove(run._r)
    para.add_run(text)


def tag_row(row, tag):
    """将 {%tr tag %} 写入该行第一个可用 TC。
    跳过 vMerge 续接格（<w:vMerge/> 无 restart val），找第一个独立格或 restart 格。
    """
    from docx.table import _Cell
    tcs = row._tr.findall(qn("w:tc"))
    target_tc = None
    for tc in tcs:
        tcPr = tc.find(qn("w:tcPr"))
        vm = tcPr.find(qn("w:vMerge")) if tcPr is not None else None
        if vm is None:
            # 无 vMerge 元素 → 独立格，直接用
            target_tc = tc
            break
        val = vm.get(qn("w:val"))
        if val == "restart":
            # vMerge restart → 竖向合并主格，可用但非首选；继续找更好的
            if target_tc is None:
                target_tc = tc
            # 不 break，继续寻找独立格
        # val is None → 续接格，跳过
    if target_tc is None:
        target_tc = tcs[0]
    c = _Cell(target_tc, row.table)
    for para in c.paragraphs:
        for run in para.runs:
            run._r.getparent().remove(run._r)
    c.paragraphs[0].add_run(tag)


def copy_row_after(table, src_row_idx):
    """在 src_row 后插入一行，内容是 src_row 的深拷贝，返回新行。"""
    src_tr = table.rows[src_row_idx]._tr
    new_tr = copy.deepcopy(src_tr)
    src_tr.addnext(new_tr)
    # docx 行列表需要刷新
    return table.rows[src_row_idx + 1]


def delete_row(table, row_idx):
    """删除指定行。"""
    tr = table.rows[row_idx]._tr
    tr.getparent().remove(tr)


# ── 模版01 — 入库申请单 ────────────────────────────────────────────────────

def build_01():
    doc = Document(os.path.join(SRC, "模版01-入库申请单.docx"))

    # 编号段落
    set_para(doc.paragraphs[1], "编号：{{ doc_number }}")

    tbl = doc.tables[0]

    # Row[0]: 项目名称(col0 标签保留), col1-2 合并值, col3 项目代号标签保留, col4 值
    set_cell(tbl.rows[0].cells[1], "{{ project.name }}")
    set_cell(tbl.rows[0].cells[4], "ZDPX-{{ project.code }}")

    # Row[1]: 申请人值, 申请日期值
    set_cell(tbl.rows[1].cells[1], "{{ project.leader }}")
    set_cell(tbl.rows[1].cells[4], "{{ activity_date }}")

    # Row[2-3]: 产品名称/标识区 — 保持"—"不变（入库单里这两行不做动态，清单在下方）

    # Row[7]: 入库√固定, 首次/变更勾选 (col3)
    set_cell(tbl.rows[7].cells[3],
             "首次  {{ '√' if is_first_checkin else '□' }}")
    # Row[8]: 变更勾选 (col3)
    set_cell(tbl.rows[8].cells[3],
             "变更  {{ '□' if is_first_checkin else '√' }}      变更报告单登记号：")

    # Row[11]: 清单示例行 → {%tr for %}
    tag_row(tbl.rows[11], "{%tr for item in items %}")
    # Row[12]: 数据行（原空行）
    set_cell(tbl.rows[12].cells[1], "{{ loop.index }}")
    set_cell(tbl.rows[12].cells[2], "{{ item.name }}")
    set_cell(tbl.rows[12].cells[4], "{{ item.code }}{{ item.version }}")
    # Row[13]: endfor
    tag_row(tbl.rows[13], "{%tr endfor %}")
    # Row[14]: 多余空行删掉
    delete_row(tbl, 14)

    # Row[17]（删一行后变16）: 底部说明（合并行，只需写第一个物理tc）
    # 行索引已变：原17->16, 18->17, 19->18（删了row14）
    set_cell(tbl.rows[16].cells[0],
             "\n{{ product_list_text }}编制完毕。通过评审。经批准，纳入受控库。")

    doc.save(os.path.join(DST, "模版01-入库申请单_tpl.docx"))
    print("✓ 模版01")


# ── 模版02 — 出库申请单 ────────────────────────────────────────────────────

def build_02():
    doc = Document(os.path.join(SRC, "模版02-出库申请单.docx"))

    set_para(doc.paragraphs[1], "编号：{{ doc_number }}")

    tbl = doc.tables[0]

    # Row[0]: 项目名称值(col1), 项目代号值(col5)
    set_cell(tbl.rows[0].cells[1], "{{ project.name }}")
    set_cell(tbl.rows[0].cells[5], "ZDPX-{{ project.code }}")

    # Row[1]: 申请人(col1), 申请日期(col5)
    set_cell(tbl.rows[1].cells[1], "{{ project.leader }}")
    set_cell(tbl.rows[1].cells[5], "{{ activity_date }}")

    # Row[4]: 出库原因勾选 — 出库√固定, 使用勾选(col2)
    set_cell(tbl.rows[4].cells[2],
             "使用 {{ '√' if checkout_type=='使用' else '□' }}")
    # Row[5]: 交付勾选(col2)
    set_cell(tbl.rows[5].cells[2],
             "交付 {{ '√' if checkout_type=='交付' else '□' }}    （交付给项目组内审）")
    # Row[6]: 变更勾选(col2)
    set_cell(tbl.rows[6].cells[2],
             "变更 {{ '√' if checkout_type=='变更' else '□' }}     变更申请单登记号：")

    # Row[10]: 第一条数据行 → {%tr for %}
    tag_row(tbl.rows[10], "{%tr for item in items %}")
    # Row[10] 现在是 for 控制行，数据写在 col1-col4
    # 用 Row[11] 作为数据行（原第2条数据）
    set_cell(tbl.rows[11].cells[1], "{{ loop.index }}")
    set_cell(tbl.rows[11].cells[2], "{{ item.name }}")
    set_cell(tbl.rows[11].cells[4], "{{ item.code }}{{ item.version }}")

    # Row[11] 后插入 endfor 行（拷贝Row[11]结构）
    endfor_row = copy_row_after(tbl, 11)
    tag_row(endfor_row, "{%tr endfor %}")

    # Row[14]: 底部文字（现在行号偏移+1变成15，但表格刷新后重新定位）
    # 找最后一行（底部说明行，原Row[14]，现在+1=15）
    last_row = tbl.rows[-1]
    set_cell(last_row.cells[0],
             "\n{{ checkout_reason_text }}经批准，将{{ product_list_text }}从受控库中检出。")

    doc.save(os.path.join(DST, "模版02-出库申请单_tpl.docx"))
    print("✓ 模版02")


# ── 模版03 — 基线发布单 ────────────────────────────────────────────────────

def build_03():
    doc = Document(os.path.join(SRC, "模版03-基线发布单.docx"))

    set_para(doc.paragraphs[1], "编号：{{ doc_number }}")

    tbl = doc.tables[0]

    # Row[0]: 项目名称(col2), 项目代号(col5)
    set_cell(tbl.rows[0].cells[2], "{{ project.name }}")
    set_cell(tbl.rows[0].cells[5], "ZDPX-{{ project.code }}")

    # Row[1]: 基线名称及标识(col2)
    set_cell(tbl.rows[1].cells[2],
             "{{ baseline.name }} ZDPX-{{ project.code }}-{{ baseline.code }} {{ baseline.version }}")

    # Row[2]: 基线版本(col2), 前向基线版本(col5)
    set_cell(tbl.rows[2].cells[2], "{{ baseline.version }}")
    set_cell(tbl.rows[2].cells[5], "{{ baseline.previous_version or '——' }}")

    # Row[3]: 发布时间(col5)
    set_cell(tbl.rows[3].cells[5], "{{ baseline.release_date }}")

    # Row[4]: 发布类型（整行合并）
    set_cell(tbl.rows[4].cells[0],
             "发布类型： 首次发布    {{ '√' if not baseline.previous_version else '□' }}"
             "        变更    {{ '□' if not baseline.previous_version else '√' }}")

    # Row[7]: 数据示例行 → {%tr for %}
    tag_row(tbl.rows[7], "{%tr for item in items %}")
    # 数据写在 for 控制行的各列（col0序号, col1名称, col3标识版本, col6备注）
    # 实际需要一行数据行，拷贝 Row[7] 后插入
    data_row = copy_row_after(tbl, 7)
    set_cell(data_row.cells[0], "{{ loop.index }}")
    set_cell(data_row.cells[1], "{{ item.name }}")
    set_cell(data_row.cells[3], "{{ item.code }} {{ item.version }}")
    set_cell(data_row.cells[6], "")

    # 在 data_row 后插入 endfor
    endfor_row = copy_row_after(tbl, 8)
    tag_row(endfor_row, "{%tr endfor %}")

    doc.save(os.path.join(DST, "模版03-基线发布单_tpl.docx"))
    print("✓ 模版03")


# ── 模版04 — 配置审核报告（产品级） ─────────────────────────────────────────

def build_04():
    doc = Document(os.path.join(SRC, "模版04-配置审核报告（产品级）.docx"))

    set_para(doc.paragraphs[1], "编号：{{ doc_number }}")

    tbl = doc.tables[0]

    # Row[0]: 项目名称(col2), 项目代号(col8)
    set_cell(tbl.rows[0].cells[2], "{{ project.name }}")
    set_cell(tbl.rows[0].cells[8], "ZDPX-{{ project.code }}")

    # Row[1]: 审核事由-入受控库勾选(col2), 入库单编号(col7)
    set_cell(tbl.rows[1].cells[2],
             "入受控库    {{ '√' if checkin_reason=='入受控库' else '□' }}")
    set_cell(tbl.rows[1].cells[7], "{{ checkin_doc_number }}")

    # Row[2]: 基线发布勾选(col2), 基线发布单编号(col7)
    set_cell(tbl.rows[2].cells[2],
             "基线发布  {{ '√' if checkin_reason=='基线发布' else '□' }}")
    set_cell(tbl.rows[2].cells[7], "{{ baseline_doc_number or '' }}")

    # Row[3]: 基线变更勾选(col2), 变更申请单编号留空(col7)
    set_cell(tbl.rows[3].cells[2],
             "基线变更  {{ '√' if checkin_reason=='基线变更' else '□' }}")
    set_cell(tbl.rows[3].cells[7], "")

    # Row[4]: 产品入库勾选(col2)
    set_cell(tbl.rows[4].cells[2],
             "产品入库  {{ '√' if checkin_reason=='产品入库' else '□' }}")
    set_cell(tbl.rows[4].cells[7], "")

    # Row[5]: 审核类型（物理√功能√ 固定，配置管理□ 固定）—— 保持原样不变

    # Row[6]: 报告日期(col2), 参加审核人员(col8)
    set_cell(tbl.rows[6].cells[2], "{{ activity_date }}")
    set_cell(tbl.rows[6].cells[8], "{{ project.cm_manager }}")

    # Row[8]: 审核内容数据行 → {%tr for %} (基线列col1, 配置项列col5)
    # 原Row[8]是单条示例数据行
    tag_row(tbl.rows[8], "{%tr for item in products %}")
    data_row = copy_row_after(tbl, 8)
    set_cell(data_row.cells[1], "{{ item.baseline_info }}")
    set_cell(data_row.cells[5],
             "{{ item.name }}\n{{ item.code }} {{ item.version }}")
    set_cell(data_row.cells[9], "")
    endfor_row = copy_row_after(tbl, 9)
    tag_row(endfor_row, "{%tr endfor %}")

    # 审核条目勾选（10条，行号因插入2行偏移+2）
    # 原Row[11-20] → 现在Row[13-22]
    offset = 2
    # 10条条目的勾选列: col5=是, col6=否, col7=不适用
    # 按 check_item(n) 条件渲染
    checklist = {
        # row_idx(原): (是条件表达式, 否条件, 不适用条件)
        11: ("check[0]=='是'", "check[0]=='否'", "check[0]=='不适用'"),
        12: ("check[1]=='是'", "check[1]=='否'", "check[1]=='不适用'"),
        13: ("check[2]=='是'", "check[2]=='否'", "check[2]=='不适用'"),
        14: ("check[3]=='是'", "check[3]=='否'", "check[3]=='不适用'"),
        15: ("check[4]=='是'", "check[4]=='否'", "check[4]=='不适用'"),
        16: ("check[5]=='是'", "check[5]=='否'", "check[5]=='不适用'"),
        17: ("check[6]=='是'", "check[6]=='否'", "check[6]=='不适用'"),
        18: ("check[7]=='是'", "check[7]=='否'", "check[7]=='不适用'"),
        19: ("check[8]=='是'", "check[8]=='否'", "check[8]=='不适用'"),
        20: ("check[9]=='是'", "check[9]=='否'", "check[9]=='不适用'"),
    }
    for orig_idx, (yes_cond, no_cond, na_cond) in checklist.items():
        row = tbl.rows[orig_idx + offset]
        set_cell(row.cells[5], "{{ '√' if " + yes_cond + " else '' }}")
        set_cell(row.cells[6], "{{ '√' if " + no_cond + " else '' }}")
        set_cell(row.cells[7], "{{ '√' if " + na_cond + " else '' }}")

    doc.save(os.path.join(DST, "模版04-配置审核报告（产品级）_tpl.docx"))
    print("✓ 模版04")


# ── 模版05 — 配置审核报告（阶段级） ─────────────────────────────────────────

def build_05():
    doc = Document(os.path.join(SRC, "模版05-配置审核报告（阶段级）.docx"))

    set_para(doc.paragraphs[1], "编号：{{ doc_number }}")

    tbl = doc.tables[0]

    # Row[0]: 项目名称(col2), 项目代号(col6)
    set_cell(tbl.rows[0].cells[2], "{{ project.name }}")
    set_cell(tbl.rows[0].cells[6], "ZDPX-{{ project.code }}")

    # Row[6]: 报告日期(col2), 参加审核人员(col7)
    set_cell(tbl.rows[6].cells[2], "{{ activity_date }}")
    set_cell(tbl.rows[6].cells[7], "{{ project.cm_manager }}")

    # 其余行（审核事由全□、审核内容全——、8条固定文本、审核结论固定）保持原样
    doc.save(os.path.join(DST, "模版05-配置审核报告（阶段级）_tpl.docx"))
    print("✓ 模版05")


# ── 模版06 — 配管清单 ─────────────────────────────────────────────────────

def build_06():
    doc = Document(os.path.join(SRC, "模版06-配管清单.docx"))

    set_para(doc.paragraphs[1], "编号：{{ doc_number }}")

    tbl = doc.tables[0]

    # Row[0]: 项目名称(col2), 项目代号(col7)
    set_cell(tbl.rows[0].cells[2], "{{ project.name }}")
    set_cell(tbl.rows[0].cells[7], "ZDPX-{{ project.code }}")

    # Row[1]: 配置管理员(col2), 记录时间(col7)
    set_cell(tbl.rows[1].cells[2], "{{ project.cm_manager }}")
    set_cell(tbl.rows[1].cells[7], "{{ activity_date }}")

    # Row[4]: 第一条数据行 → {%tr for %}
    tag_row(tbl.rows[4], "{%tr for item in products %}")

    # Row[5]: 第二条数据行 → 实际数据模板行
    set_cell(tbl.rows[5].cells[0], "{{ loop.index }}")
    set_cell(tbl.rows[5].cells[1], "{{ item.name }}")
    set_cell(tbl.rows[5].cells[3], "{{ item.code }}")
    set_cell(tbl.rows[5].cells[4], "{{ item.version }}")
    set_cell(tbl.rows[5].cells[5], "{{ item.previous_version or '——' }}")
    set_cell(tbl.rows[5].cells[6],
             "{{ item.baseline_name + ' ' + item.baseline_version if item.baseline_id else '' }}")
    set_cell(tbl.rows[5].cells[7], "{{ '√' if item.status=='有效' else '' }}")
    set_cell(tbl.rows[5].cells[8], "")
    set_cell(tbl.rows[5].cells[9], "{{ '√' if item.status=='无效' else '' }}")

    # 在 Row[5] 后插入 endfor
    endfor_row = copy_row_after(tbl, 5)
    tag_row(endfor_row, "{%tr endfor %}")

    doc.save(os.path.join(DST, "模版06-配管清单_tpl.docx"))
    print("✓ 模版06")


# ── 模版07 — 配管报告 ─────────────────────────────────────────────────────

def build_07():
    """
    原始模板行结构（0-indexed）：
      0: 项目名称
      1: 当前阶段/报告时间范围/配置管理员
      2: 变更情况 header
      3: 序号/日期/变更申请单/状态 sub-header
      4: 1 (empty变更数据行)
      5: 基线建立情况 header
      6: 序号/基线名称/标识/发布日期/状态 sub-header
      7: 1. 分配基线 ... (唯一示例数据行)
      8: 配置审核情况 header
      9: 序号/配置项名称/审核时间/发现问题/处理情况 sub-header
     10: 1. 开发计划 ... (示例数据行1)
     11: 2. 分配基线 ... (示例数据行2)
     12: 3. 配置项测试计划 ... (示例数据行3)
     13: 当前有效基线 header
     14: 序号/基线名称/基线版本/配置项名称/标识版本 sub-header
     15: 1. 功能基线 ... (示例数据行1)
     16: 2. 分配基线 ... (示例数据行2)
     17: 当前配置项清单 header
     18: 序号/配置项名称/标识/版本/状态 sub-header
     19: 1. ... (示例数据行1)
     20: 2. ... (示例数据行2)
     21: 3. ... (示例数据行3)
    策略：从下往上处理，先删多余示例行再 tag+insert，避免行号偏移。
    """
    doc = Document(os.path.join(SRC, "模版07-配管报告.docx"))
    set_para(doc.paragraphs[1], "编号：{{ doc_number }}")
    tbl = doc.tables[0]

    # 静态字段
    set_cell(tbl.rows[0].cells[2], "{{ project.name }}")
    set_cell(tbl.rows[0].cells[9], "ZDPX-{{ project.code }}")
    set_cell(tbl.rows[1].cells[2], "{{ phase.name }}")
    set_cell(tbl.rows[1].cells[8], "{{ phase.start_date }} {{ activity_date }}")
    set_cell(tbl.rows[1].cells[14], "{{ project.cm_manager }}")

    # ── Section 4: 当前配置项清单 (原 Row[17-21]) ──────────────────────────
    # 删除示例行2、3（从下往上）
    delete_row(tbl, 21)
    delete_row(tbl, 20)
    # 现在 Row[19] 是唯一数据行
    tag_row(tbl.rows[19], "{%tr for item in all_products %}")
    ci_data = copy_row_after(tbl, 19)
    set_cell(ci_data.cells[0], "{{ loop.index }}")
    set_cell(ci_data.cells[1], "{{ item.name }}")
    set_cell(ci_data.cells[6], "{{ item.code }}")
    set_cell(ci_data.cells[12], "{{ item.version }}")
    set_cell(ci_data.cells[13], "{{ item.status }}")
    ci_endfor = copy_row_after(tbl, 20)
    tag_row(ci_endfor, "{%tr endfor %}")
    # 现在：19=for, 20=data, 21=endfor

    # ── Section 3: 当前有效基线 (原 Row[13-16]) ────────────────────────────
    # 删除示例行2
    delete_row(tbl, 16)
    # 现在 Row[15] 是唯一数据行
    tag_row(tbl.rows[15], "{%tr for row in baseline_product_rows %}")
    bpr_data = copy_row_after(tbl, 15)
    set_cell(bpr_data.cells[0], "{{ loop.index }}")
    set_cell(bpr_data.cells[1], "{{ row.baseline_name }}")
    set_cell(bpr_data.cells[3], "{{ row.baseline_version }}")
    set_cell(bpr_data.cells[4], "{{ row.product_name }}")
    set_cell(bpr_data.cells[10], "{{ row.product_code }} {{ row.product_version }}")
    bpr_endfor = copy_row_after(tbl, 16)
    tag_row(bpr_endfor, "{%tr endfor %}")
    # 现在：15=for, 16=data, 17=endfor

    # ── Section 2: 配置审核情况 (原 Row[8-12]) ─────────────────────────────
    # 删除示例行2、3
    delete_row(tbl, 12)
    delete_row(tbl, 11)
    # 现在 Row[10] 是唯一数据行
    tag_row(tbl.rows[10], "{%tr for ar in phase_audit_reports %}")
    ar_data = copy_row_after(tbl, 10)
    set_cell(ar_data.cells[0], "{{ loop.index }}")
    set_cell(ar_data.cells[1], "{{ ar.product_names }}")
    set_cell(ar_data.cells[6], "{{ ar.activity_date }}")
    set_cell(ar_data.cells[8], "")
    set_cell(ar_data.cells[10], "")
    ar_endfor = copy_row_after(tbl, 11)
    tag_row(ar_endfor, "{%tr endfor %}")
    # 现在：10=for, 11=data, 12=endfor

    # ── Section 1: 基线建立情况 (原 Row[5-7]) ──────────────────────────────
    # Row[7] 唯一示例数据行
    tag_row(tbl.rows[7], "{%tr for bl in phase_baselines %}")
    bl_data = copy_row_after(tbl, 7)
    set_cell(bl_data.cells[0], "{{ loop.index }}")
    set_cell(bl_data.cells[1], "{{ bl.name }}")
    set_cell(bl_data.cells[6],
             "ZDPX-{{ project.code }}-{{ bl.code }} {{ bl.version }}")
    set_cell(bl_data.cells[10], "{{ bl.release_date }}")
    set_cell(bl_data.cells[13], "{{ bl.status }}")
    bl_endfor = copy_row_after(tbl, 8)
    tag_row(bl_endfor, "{%tr endfor %}")
    # 现在：7=for, 8=data, 9=endfor

    doc.save(os.path.join(DST, "模版07-配管报告_tpl.docx"))
    print("✓ 模版07")


# ── 主入口 ────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    print("开始生成模板...")
    build_01()
    build_02()
    build_03()
    build_04()
    build_05()
    build_06()
    build_07()
    print(f"\n完成！输出目录：{DST}")
