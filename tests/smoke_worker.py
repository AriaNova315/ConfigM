"""Internal worker: requires the disposable directory created by test_deployment."""
import asyncio
import io
import json
from pathlib import Path
import sys
import zipfile
from unittest.mock import patch

assert (Path.cwd() / ".deployment-test").is_file(), "Only run in a disposable verification copy"
assert not (Path.cwd() / "configm.db").exists(), "Verification requires a fresh database"
sys.path.insert(0, str(Path.cwd()))

from openpyxl import load_workbook
from app.main import app


async def request(method, url, data=None, content_type="application/json"):
    body = data if isinstance(data, bytes) else json.dumps(data).encode() if data is not None else b""
    sent = []
    consumed = False

    async def receive():
        nonlocal consumed
        if not consumed:
            consumed = True
            return {"type": "http.request", "body": body, "more_body": False}
        # Keep streaming responses alive until the application finishes sending.
        await asyncio.Event().wait()

    async def send(message):
        sent.append(message)

    scope = {"type": "http", "asgi": {"version": "3.0", "spec_version": "2.4"},
             "http_version": "1.1", "method": method, "scheme": "http", "path": url,
             "raw_path": url.encode(), "query_string": b"", "root_path": "",
             "headers": [(b"content-type", content_type.encode()), (b"content-length", str(len(body)).encode())],
             "server": ("127.0.0.1", 8000), "client": ("127.0.0.1", 12345)}
    await app(scope, receive, send)
    status = next(m["status"] for m in sent if m["type"] == "http.response.start")
    content = b"".join(m.get("body", b"") for m in sent if m["type"] == "http.response.body")
    return status, content


async def api(method, url, data=None):
    status, body = await request(method, url, data)
    assert status < 400, (method, url, status, body)
    return json.loads(body) if body else None


async def main():
    async with app.router.lifespan_context(app):
        assert await api("GET", "/api/projects") == []
        assert len(await api("GET", "/api/activity-types")) == 5
    async with app.router.lifespan_context(app):
        assert len(await api("GET", "/api/activity-types")) == 5
        assert (await request("GET", "/docs"))[0] == 404
        assert (await request("GET", "/redoc"))[0] == 404
        html = (await request("GET", "/"))[1].decode()
        assert "https://cdn" not in html and "/static/vendor/alpine-3.15.12.min.js" in html
        status, js = await request("GET", "/static/vendor/alpine-3.15.12.min.js")
        assert status == 200 and len(js) > 10000

        workbook = load_workbook("resources/ConfigM导入模板.xlsx")
        assert workbook.sheetnames == ["项目信息", "项目进度计划", "配置项清单", "基线组成"]
        assert all(not any(row) for ws in workbook for row in ws.iter_rows(min_row=2, values_only=True))
        plan = workbook["项目进度计划"]
        # Styled empty rows must be filled at explicit row numbers.
        rows = [[1, "CON:设计?阶段", "2026-01-01", "2026-01-31", None],
                ["1.1", "配置管理活动", None, None, None]]
        for i, task in enumerate(["产品入库", "发布功能基线", "产品出库", "配置状态纪实", "配置管理报告"], 1):
            rows.append([f"1.1.{i}", task, f"2026-01-{i:02d}", f"2026-01-{i:02d}", "测试文档V1.0" if "库" in task else None])
        for i, row in enumerate(rows, 2):
            for j, value in enumerate(row, 1):
                plan.cell(i, j, value)
        for i, value in enumerate([1, "测试文档", "DOC-001", "V1.0"], 1):
            workbook["配置项清单"].cell(2, i, value)
        for i, value in enumerate(["功能基线", "BL-001", "V1.0", "测试文档", "DOC-001", "V1.0"], 1):
            workbook["基线组成"].cell(2, i, value)
        # Fill explicit rows in the styled blank template.
        for i, value in enumerate(["离线验证:项目?", "00001", "测试组长", "测试管理员"], 1):
            workbook["项目信息"].cell(2, i, value)
        buf = io.BytesIO()
        workbook.save(buf)
        workbook.close()
        boundary = "ConfigMVerificationBoundary"
        upload = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"synthetic.xlsx\"\r\n"
                  "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n").encode() + buf.getvalue() + f"\r\n--{boundary}--\r\n".encode()
        status, body = await request("POST", "/api/import/parse", upload, f"multipart/form-data; boundary={boundary}")
        assert status == 200, body
        parsed = json.loads(body)
        assert parsed["summary"]["activities"] == 5
        result = await api("POST", "/api/import/confirm", {
            "checkin_reasons": {x["key"]: "入受控库" for x in parsed["needs_checkin"]},
            "checkout_reasons": {x["key"]: "离线验收" for x in parsed["needs_checkout"]},
            "baseline_selections": {x["key"]: "功能基线" for x in parsed["needs_baseline"]}})
        projects = await api("GET", "/api/projects")
        assert len(projects) == 1 and projects[0]["name"] == "离线验证:项目?"
        pid = projects[0]["id"]
        await api("POST", f"/api/projects/{pid}/generate-all")
        records = await api("GET", f"/api/projects/{pid}/documents")
        assert len(records) == 7, records
        assert all(r["generated_at"] for r in records)
        with patch("app.core.output_paths.os.replace", side_effect=PermissionError("locked")):
            status, error = await request("POST", f"/api/documents/{records[0]['id']}/generate")
            assert status == 409 and "关闭" in json.loads(error)["detail"]
        unchanged = await api("GET", f"/api/projects/{pid}/documents")
        assert unchanged[0]["generated_at"] == records[0]["generated_at"]
        for record in records:
            status, content = await request("GET", f"/api/documents/{record['id']}/download")
            assert status == 200
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                assert z.testzip() is None
                assert b"{{" not in z.read("word/document.xml")
            assert Path(record["file_path"]).resolve().is_relative_to(Path("generated").resolve())
        assert (await request("GET", f"/api/projects/{pid}/download-all"))[0] == 200
        status, exported = await request("GET", "/api/export")
        assert status == 200
        from app.services.importer import parse_excel
        data = parse_excel(exported)
        assert data["project"]["name"] == projects[0]["name"]
        assert len(data["products"]) == 1 and len(data["baselines"]) == 1
        await api("PUT", f"/api/projects/{pid}", {"leader": "修改后的测试组长"})
    async with app.router.lifespan_context(app):
        projects = await api("GET", "/api/projects")
        assert projects[0]["leader"] == "修改后的测试组长"
        assert len(await api("GET", "/api/activity-types")) == 5
    print("PASS: fresh/repeated startup, local resources, Excel parse/confirm/export, seven DOCX types, downloads and persistence")


asyncio.run(main())
