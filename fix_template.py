from docx import Document

def merge_runs(paragraph):
    if len(paragraph.runs) <= 1:
        return
    full_text = paragraph.text
    paragraph.runs[0].text = full_text
    for run in paragraph.runs[1:]:
        run._element.getparent().remove(run._element)

doc = Document("test_template.docx")
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                if "{%" in para.text or "{{" in para.text:
                    merge_runs(para)

doc.save("test_template_fixed.docx")
print("修复完成：test_template_fixed.docx")
