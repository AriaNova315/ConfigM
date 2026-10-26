from docx import Document
import lxml.etree as etree

doc = Document("test_template_fixed.docx")
for i, table in enumerate(doc.tables):
    print(f"=== 表格 {i} ===")
    for j, row in enumerate(table.rows):
        print(f"  行 {j}:")
        for k, cell in enumerate(row.cells):
            runs = cell.paragraphs[0].runs
            print(f"    列 {k}: runs={[r.text for r in runs]!r}  full={cell.text!r}")

# 打印第2行第1列的原始XML，看有没有残留元素
print("\n=== 行1列0 原始XML ===")
row1 = doc.tables[0].rows[1]
print(etree.tostring(row1.cells[0]._element, pretty_print=True).decode())
