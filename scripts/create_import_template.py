"""Build the empty, value-only workbook accepted by the existing Excel importer."""
from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SHEETS = {
    "项目信息": ["项目名称", "项目代号", "项目组长", "配置管理员"],
    "项目进度计划": ["WBS", "任务名称", "计划开始时间", "计划完成时间", "关联配置项"],
    "配置项清单": ["序号", "配置项名称", "配置项标识", "版本"],
    "基线组成": ["基线名称", "基线标识", "基线版本", "包含配置项名称", "配置项标识", "配置项版本"],
}


def build(path: Path):
    wb = Workbook()
    wb.remove(wb.active)
    for name, headings in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(headings)
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 32
        for col, heading in enumerate(headings, 1):
            cell = ws.cell(1, col)
            cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill("solid", fgColor="1A3A5C")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = max(18, len(heading) * 2 + 4)
        # Text inputs retain leading zeroes and WBS such as 1.1.1.
        for row in ws.iter_rows(min_row=2, max_row=101, max_col=len(headings)):
            for cell in row:
                cell.number_format = "@"
        ws.sheet_view.showGridLines = True
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "resources" / "ConfigM导入模板.xlsx"
    if target.exists():
        raise SystemExit(f"Refusing to overwrite existing file: {target}")
    build(target)
